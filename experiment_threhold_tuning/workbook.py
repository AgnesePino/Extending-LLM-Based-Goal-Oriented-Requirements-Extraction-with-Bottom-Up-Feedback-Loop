"""Self-contained, human-annotated threshold workbook. No model inference."""
import csv
import hashlib
import json
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

FILENAME = 'threshold_tuning_manual_review.xlsx'
PAIR_SHEETS = {'prime_ground_truth': 'Confronti ground truth',
               'prime_parent': 'Confronti padre',
               'prime_prime_duplicate': 'Confronti duplicati'}
FIELDS = ['pair_id', 'dataset', 'split', 'pair_type', 'left_id', 'right_id',
          'left_text', 'right_text', 'score', 'equivalent', 'notes']


def evaluate_threshold(records: list[dict], threshold: float) -> dict:
    tp = fp = tn = fn = 0
    for record in records:
        predicted = record["score"] >= threshold
        actual = record["equivalent"]
        tp += int(predicted and actual)
        fp += int(predicted and not actual)
        tn += int(not predicted and not actual)
        fn += int(not predicted and actual)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    specificity = tn / (tn + fp) if tn + fp else 0.0
    return {"threshold": threshold, "tp": tp, "fp": fp, "tn": tn, "fn": fn,
            "precision": precision, "recall": recall,
            "f1": 2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else 0.0,
            "balanced_accuracy": (recall + specificity) / 2,
            "margin": min(abs(record["score"] - threshold) for record in records)}


def calibrate_parent_threshold(records: list[dict]) -> dict:
    """Maximize equivalent-class F1 on a 0.01 grid.

    Break ties by balanced accuracy, then distance from the nearest observed
    score. The grid and margin avoid a brittle cut between nearly equal scores.
    This policy favors recovering paraphrases while reporting false matches.
    """
    if not records or {record.get("equivalent") for record in records} != {True, False}:
        raise ValueError("Calibration requires both equivalent and non-equivalent pairs")
    for record in records:
        if type(record["equivalent"]) is not bool:
            raise ValueError("Equivalence annotations must be booleans")
        score = record.get("score")
        if (isinstance(score, bool) or not isinstance(score, (int, float))
                or not math.isfinite(score) or not -1 <= score <= 1):
            raise ValueError("Calibration requires finite cosine scores in [-1, 1]")
    grid = [evaluate_threshold(records, round(i / 100, 2)) for i in range(101)]
    best = max(grid, key=lambda row: (row["f1"], row["balanced_accuracy"], row["margin"]))
    return {"selected_threshold": best["threshold"], "metrics": best,
            "selection": "max F1; ties: balanced accuracy, then margin; grid step 0.01",
            "pair_count": len(records), "equivalent_count": sum(r["equivalent"] for r in records),
            "held_out_pair_count": 0, "validated": False, "grid": grid}


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def _hash(value):
    return hashlib.sha256(value.encode('utf-8')).hexdigest()


def _table(book, name, headers, rows, editable=()):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        if any(isinstance(value, str) and len(value) > 32767 for value in row):
            raise ValueError(f'Text exceeds the Excel cell limit in {name}; cannot truncate source text')
        sheet.append([None if value == '' else value for value in row])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                # Source descriptions must remain literal, including leading '='.
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.protection = Protection(locked=cell.row == 1 or headers[cell.column-1] not in editable)
            if cell.row == 1:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='24476A')
            elif headers[cell.column-1] in editable:
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
    for col in sheet.columns:
        header = col[0].value
        width = 70 if header in ('left_text', 'right_text', 'description', 'hlg_prime', 'text', 'parent_description') else 25
        if header in ('notes', 'reason', 'istruzioni'):
            width = 65
        sheet.column_dimensions[col[0].column_letter].width = width
    for row in range(2, sheet.max_row + 1):
        longest = max(len(str(c.value or '')) for c in sheet[row])
        sheet.row_dimensions[row].height = min(300, max(32, 15 * (longest // 65 + 1)))
    sheet.freeze_panes = 'C2'
    if rows:
        table = Table(displayName=f'Table{len(book.worksheets)}', ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = False
    sheet.sheet_view.zoomScale = 80
    return sheet


def _rows(sheet):
    data = list(sheet.values)
    return [dict(zip(data[0], row)) for row in data[1:] if any(v is not None for v in row)]


def write_workbook(path, manifest, coverage, snapshots, pairs, annotations=None):
    """One output file containing readable tables and sealed source snapshots."""
    path = Path(path)
    if path.exists():
        raise FileExistsError(f'Workbook already exists; annotations preserved: {path}')
    book = Workbook()
    book.remove(book.active)
    _table(book, 'Guida', ['argomento', 'istruzioni'], [
        ['Scopo', 'Tutti i confronti disponibili, senza filtri sulla cosine. Nessun giudizio automatico.'],
        ['Annotazione', 'Nei tre fogli Confronti compilare solo le celle gialle: equivalent=true/false e notes. Vuoto significa non ancora annotato.'],
        ['Equivalenza', 'Giudicare lo stesso risultato funzionale; il punteggio non dimostra equivalenza o supporto documentale.'],
        ['Dati', 'LLG, HLG originali, HLG ricostruiti e ground truth sono riportati integralmente nei rispettivi fogli.'],
        ['Metriche', 'Eseguire tune sul workbook dopo le annotazioni per aggiornare Griglia soglie e Riepilogo tuning.'],
        ['Scelta finale', 'Nel foglio Scelta soglie inserire un solo valore in B2 e confirmed=true in C2. La riga duplicati riusa la stessa soglia. La proposta automatica non è una conferma.'],
        ['Validazione', 'Le soglie candidate usano solo calibration; validation serve a misurare il valore proposto o confermato.'],
        ['Ground truth', 'Il confronto HLG prime–ground truth è diagnostico e non sostituisce la calibrazione padre–ricostruzione.'],
        ['Integrità', 'I fogli protetti prevengono modifiche accidentali. I dati congelati sono incorporati; non sono una firma contro manomissioni intenzionali.'],
        ['Copertura', 'I dataset esclusi sono visibili ma non hanno cosine inventate. La campagna richiede 12 baseline; la copertura del tuning è dichiarata separatamente.'],
        ['Score', 'Le cosine sono memorizzate come testo decimale per conservare esattamente la precisione originale.'],
    ])
    _table(book, 'Copertura dataset', ['dataset', 'status', 'reason', 'group_count'],
           [[r.get(k) for k in ('dataset', 'status', 'reason', 'group_count')] for r in coverage])
    children, parents, primes, references, provenance = [], [], [], [], []
    for source in snapshots:
        name = source['dataset']
        for i, goal in enumerate(source['structuredHighLevelGoals']['goals']):
            parents.append([name, i, goal['name'], goal['description'], 'top_down_collection'])
        for group in source['reconstructions']:
            parent = group['parent']
            primes.append([name, group['group_id'], parent['name'], parent['description'], group['hlg_prime']])
            for i, child in enumerate(group['llgs']):
                children.append([name, group['group_id'], i, child['name'], child['description'],
                                 child['high_level_associated']['name'], child['high_level_associated']['description']])
        for i, text in enumerate(source['ground_truth_hlg']):
            references.append([name, f'gt_{i:03d}', text])
        provenance.append([name, 'source_provenance', _json(source['provenance'])])
    for key, value in manifest.items():
        provenance.append(['RUN', key, _json(value)])
    _table(book, 'LLG origine', ['dataset', 'group_id', 'index', 'name', 'description', 'parent_name', 'parent_description'], children)
    _table(book, 'HLG originali', ['dataset', 'index', 'name', 'description', 'source'], parents)
    _table(book, 'HLG ricostruiti', ['dataset', 'group_id', 'parent_name', 'description', 'hlg_prime'], primes)
    _table(book, 'Ground truth', ['dataset', 'reference_id', 'text'], references)
    annotations = annotations or {}
    for kind, sheet_name in PAIR_SHEETS.items():
        rows = []
        for pair in pairs:
            if pair['pair_type'] != kind:
                continue
            annotation = annotations.get(pair['pair_id'], {})
            rows.append([str(pair[k]) if k == 'score' else pair[k] for k in FIELDS[:9]] +
                        [annotation.get('equivalent') or None, annotation.get('notes') or None])
        sheet = _table(book, sheet_name, FIELDS, rows, editable=('equivalent', 'notes'))
        rule = DataValidation(type='list', formula1='"true,false"', allow_blank=True)
        rule.errorTitle, rule.error, rule.showErrorMessage = 'Giudizio non valido', 'Usare true, false o vuoto.', True
        sheet.add_data_validation(rule)
        rule.add(f'J2:J{max(2, sheet.max_row)}')
    # Wide matrices retain every reference column; descriptions are resolved in Ground truth.
    sheet = book.create_sheet('Matrici complete')
    for source in snapshots:
        name = source['dataset']
        current = [p for p in pairs if p['dataset'] == name and p['pair_type'] == 'prime_ground_truth']
        refs = [f'gt_{i:03d}' for i in range(len(source['ground_truth_hlg']))]
        sheet.append([name])
        sheet.append(['group_id', *refs])
        by_id = {(p['left_id'], p['right_id']): str(p['score']) for p in current}
        for group in source['reconstructions']:
            sheet.append([group['group_id'], *[by_id[(group['group_id'], ref)] for ref in refs]])
        sheet.append([])
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 30
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    sheet.freeze_panes = 'B3'
    sheet.protection.sheet = True
    _table(book, 'Provenienza', ['dataset', 'field', 'text'], provenance)
    selection = _table(book, 'Scelta soglie', ['pair_type', 'threshold', 'confirmed', 'notes', 'selection_source'],
           [[kind, None, None, None, 'manual_review'] for kind in ('prime_parent', 'prime_prime_duplicate')],
           editable=('threshold', 'confirmed', 'notes', 'selection_source'))
    for column in ('B', 'C', 'E'):
        cell = selection[f'{column}3']
        cell.value = f'={column}2'
        cell.protection = Protection(locked=True)
    # Seal all static visible tables, including matrices; annotation cells stay editable.
    frozen_tables = {}
    for sheet in book:
        if sheet.title != 'Scelta soglie':
            frozen_tables[sheet.title] = [list(row) for row in sheet.values]
    payload = _json(dict(manifest=manifest, coverage=coverage, snapshots=snapshots,
                         pairs=pairs, tables=frozen_tables))
    technical = book.create_sheet('_frozen')
    technical.append(['sha256', _hash(payload)])
    for offset in range(0, len(payload), 30000):
        technical.append(['chunk', payload[offset:offset+30000]])
        technical.cell(technical.max_row, 2).data_type = 's'
    technical.sheet_state = 'veryHidden'
    path.parent.mkdir(parents=True, exist_ok=True)
    _table(book, 'Riepilogo tuning', ['pair_type', 'status'],
           [[kind, 'awaiting_annotations'] for kind in PAIR_SHEETS])
    book.save(path)
    return path


def consolidate(directory):
    """Migrate an archived run without inference or recomputing a score."""
    directory = Path(directory)
    path = directory / FILENAME
    if path.exists():
        return path
    manifest = json.loads((directory/'threshold_tuning_manifest.json').read_text(encoding='utf-8'))
    raw = (directory/'frozen_cosine_pairs.json').read_bytes()
    if hashlib.sha256(raw).hexdigest() != manifest['pairs_sha256']:
        raise ValueError('Frozen pair scores changed')
    with (directory/'dataset_coverage.csv').open(encoding='utf-8-sig', newline='') as stream:
        coverage = list(csv.DictReader(stream))
    snapshots = []
    for item in manifest['matrices']:
        raw_snapshot = (directory/item['input']).read_bytes()
        if hashlib.sha256(raw_snapshot).hexdigest() != item['input_sha256']:
            raise ValueError('Frozen input snapshot changed')
        snapshots.append(json.loads(raw_snapshot))
    with (directory/'cosine_pairs_to_annotate.csv').open(encoding='utf-8-sig', newline='') as stream:
        annotations = {r['pair_id']: r for r in csv.DictReader(stream)}
    original = {p['pair_id']: p for p in json.loads(raw)}
    if set(annotations) != set(original):
        raise ValueError('Archive annotation IDs differ from frozen pairs')
    for row in annotations.values():
        p = original[row['pair_id']]
        if any(row[k] != str(p[k]) for k in FIELDS[:8]) or float(row['score']) != p['score']:
            raise ValueError('Archive annotations changed immutable fields')
        if row.get('equivalent', '').strip().lower() not in ('', 'true', 'false'):
            raise ValueError('Archive labels must be true, false or blank')
    path = write_workbook(path, manifest, coverage, snapshots, json.loads(raw), annotations)
    read_workbook(path)
    return path


def read_workbook(path):
    book = load_workbook(path)
    technical = book['_frozen']
    payload = ''.join(row[1] for row in list(technical.values)[1:])
    if _hash(payload) != technical['B1'].value:
        raise ValueError('Frozen workbook data changed')
    frozen = json.loads(payload)
    for title, expected in frozen['tables'].items():
        actual = [list(row) for row in book[title].values]
        if title in PAIR_SHEETS.values():
            # Allow sorting/filtering the rows; only the last two columns are editable.
            actual = sorted([r[:9] for r in actual[1:]], key=lambda r: str(r[0]))
            expected = sorted([r[:9] for r in expected[1:]], key=lambda r: str(r[0]))
        if actual != expected:
            raise ValueError(f'Immutable texts, scores or rows changed in {title}')
    labeled = []
    frozen_pairs = {p['pair_id']: p for p in frozen['pairs']}
    for title in PAIR_SHEETS.values():
        for row in _rows(book[title]):
            value = row.get('equivalent')
            label = '' if value is None else str(value).strip().lower()
            if label not in ('', 'true', 'false'):
                raise ValueError('Use true, false or blank for equivalence')
            if label:
                labeled.append(dict(frozen_pairs[row['pair_id']], equivalent=label == 'true'))
    return book, frozen, labeled


def _selection_rows(book):
    rows = _rows(book['Scelta soglie'])
    if [r['pair_type'] for r in rows] != ['prime_parent', 'prime_prime_duplicate']:
        raise ValueError('Preserve the two threshold selection rows')
    for field, column in (('threshold', 'B'), ('confirmed', 'C'), ('selection_source', 'E')):
        if rows[1].get(field) == f'={column}2':
            rows[1][field] = rows[0].get(field)
    return rows


def _selection(book):
    rows = _selection_rows(book)
    result = {}
    for row in rows:
        confirmed = str(row['confirmed'] or '').strip().lower()
        if confirmed not in ('', 'true', 'false'):
            raise ValueError('Confirmation must be true, false or blank')
        if confirmed == 'true':
            value = row['threshold']
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 <= value <= 1:
                raise ValueError('A confirmed threshold must be a number in [0, 1]')
            result[row['pair_type']] = float(value)
    if len(set(result.values())) > 1:
        raise ValueError('Equivalence and duplicate candidates require one shared similarity threshold')
    return result


def analyze_workbook(book, frozen, labeled):
    choices, results, grid_rows = _selection(book), {}, []
    sources = {row['pair_type']: row.get('selection_source') or 'manual_review'
               for row in _selection_rows(book)}
    if any(value not in {'manual_review', 'delegated_provisional'} for value in sources.values()):
        raise ValueError('selection_source must be manual_review or delegated_provisional')
    for kind in PAIR_SHEETS:
        subset = [p for p in labeled if p['pair_type'] == kind]
        calibration = [p for p in subset if p['split'] == 'calibration']
        validation = [p for p in subset if p['split'] == 'validation']
        total = sum(p['pair_type'] == kind for p in frozen['pairs'])
        info = dict(status='awaiting_both_calibration_classes', total=total, labeled=len(subset),
                    calibration_count=len(calibration), validation_count=len(validation))
        if {p['equivalent'] for p in calibration} == {True, False}:
            proposal = calibrate_parent_threshold(calibration)
            grid_rows.extend(dict(pair_type=kind, **r) for r in proposal.pop('grid'))
            proposed = proposal['selected_threshold']
            selected = choices.get(kind)
            measured = proposed if selected is None else selected
            delegated = sources.get(kind) == 'delegated_provisional'
            info.update(status=('user_directed_provisional' if delegated else 'manually_confirmed')
                        if selected is not None else 'awaiting_manual_confirmation',
                        proposed_threshold=proposed, confirmed_threshold=selected,
                        calibration_metrics=evaluate_threshold(calibration, measured),
                        validation_metrics=evaluate_threshold(validation, measured) if validation else None,
                        annotation_complete=len(subset) == total,
                        selection_source=sources.get(kind),
                        validated=not delegated and len(subset) == total and {p['equivalent'] for p in validation} == {True, False})
        elif kind in choices:
            if sources.get(kind) == 'delegated_provisional':
                info.update(status='user_directed_provisional', confirmed_threshold=choices[kind],
                            validated=False, annotation_complete=len(subset) == total,
                            selection_source='delegated_provisional',
                            limitation='User-directed value; no labeled calibration estimate. Not a completed human review.')
            else:
                raise ValueError(f'Cannot confirm {kind} before both calibration classes are annotated')
        results[kind] = info
    return results, grid_rows


def tune_workbook(path):
    book, frozen, labeled = read_workbook(path)
    results, grid = analyze_workbook(book, frozen, labeled)
    for title in ('Griglia soglie', 'Riepilogo tuning'):
        if title in book:
            del book[title]
    headers = ['pair_type', 'threshold', 'tp', 'fp', 'tn', 'fn', 'precision', 'recall', 'f1', 'balanced_accuracy', 'margin']
    _table(book, 'Griglia soglie', headers, [[row[k] for k in headers] for row in grid])
    headers = ['pair_type', 'status', 'total', 'labeled', 'calibration_count', 'validation_count',
               'proposed_threshold', 'confirmed_threshold', 'annotation_complete', 'validated',
               'calibration_metrics', 'validation_metrics']
    rows = [[kind, *[_json(info[k]) if isinstance(info.get(k), dict) else info.get(k) for k in headers[1:]]]
            for kind, info in results.items()]
    _table(book, 'Riepilogo tuning', headers, rows)
    book.save(path)
    return results


def confirmed_thresholds(path):
    """Freeze explicit choices and distinguish reviewed from delegated provisional values."""
    book, frozen, labeled = read_workbook(path)
    results, _ = analyze_workbook(book, frozen, labeled)
    kinds = ('prime_parent', 'prime_prime_duplicate')
    if any(results[k]['status'] not in {'manually_confirmed', 'user_directed_provisional'} for k in kinds):
        raise ValueError('Confirm both parent and duplicate thresholds in Scelta soglie before running experiments')
    return dict(similarity_threshold=results['prime_parent']['confirmed_threshold'],
                parent_threshold=results['prime_parent']['confirmed_threshold'],
                duplicate_threshold=results['prime_prime_duplicate']['confirmed_threshold'],
                workbook=str(Path(path).resolve()), workbook_sha256=hashlib.sha256(Path(path).read_bytes()).hexdigest(),
                run_id=frozen['manifest']['run_id'],
                embedding_configuration=frozen['manifest']['embedding_configuration'],
                resolved_embedding_revision=frozen['manifest'].get('resolved_embedding_revision'),
                selection_status={k: results[k]['status'] for k in kinds},
                validation={k: results[k] for k in kinds})
